#!/usr/bin/env python3
"""
Build the formatted Word document from the Markdown source.

Formatting contract (as specified for the deliverable):
  * every run is black (RGB 0,0,0)
  * nothing is bold, at any level
  * heading 18 pt, subheading 16 pt, subsection title 14 pt, body 12 pt
  * tables and code blocks are 12 pt as well

Source : MapCompete_Technical_Documentation.md
Output : MapCompete_Technical_Documentation.docx

Usage: python build_documentation_docx.py
"""

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches

ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "MapCompete_Technical_Documentation.md"
OUTPUT = ROOT / "MapCompete_Technical_Documentation.docx"

BLACK = RGBColor(0x00, 0x00, 0x00)
BODY_FONT = "Calibri"
CODE_FONT = "Consolas"
SIZE_HEADING = 18      # document title and numbered sections
SIZE_SUBHEADING = 16   # subsections
SIZE_TITLE = 14        # title block and third-level titles
SIZE_BODY = 12         # paragraphs, points, tables, code

NUMBERED_RE = re.compile(r"^\d{1,2}\.\s+(.*)$")


def strip_inline(text):
    """Remove Markdown inline markers (backticks) - the text stays plain."""
    return text.replace("`", "").strip()


def style_run(run, size, font=BODY_FONT):
    """Apply the single allowed text appearance: black, not bold, given size."""
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = False
    run.font.italic = False
    run.font.color.rgb = BLACK
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), font)


def add_heading(doc, text, size):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(12 if size >= 16 else 8)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.keep_with_next = True
    style_run(paragraph.add_run(strip_inline(text)), size)
    return paragraph


def add_body(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    style_run(paragraph.add_run(strip_inline(text)), SIZE_BODY)
    return paragraph


def add_point(doc, text, style):
    paragraph = doc.add_paragraph(style=style)
    paragraph.paragraph_format.space_after = Pt(3)
    style_run(paragraph.add_run(strip_inline(text)), SIZE_BODY)
    return paragraph


def add_code_block(doc, lines):
    for index, line in enumerate(lines):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(0 if index < len(lines) - 1 else 8)
        paragraph.paragraph_format.space_before = Pt(4 if index == 0 else 0)
        paragraph.paragraph_format.left_indent = Inches(0.2)
        style_run(paragraph.add_run(line.rstrip()), SIZE_BODY, font=CODE_FONT)


def add_table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for row_index, row in enumerate(rows):
        for column_index, cell_text in enumerate(row):
            paragraph = table.cell(row_index, column_index).paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(2)
            style_run(paragraph.add_run(strip_inline(cell_text)), SIZE_BODY)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return table


def configure_styles(doc):
    """Base style plus explicit heading styles, all black and never bold."""
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(SIZE_BODY)
    normal.font.color.rgb = BLACK
    normal.font.bold = False

    for name, size in (("Heading 1", SIZE_HEADING), ("Heading 2", SIZE_SUBHEADING),
                       ("Heading 3", SIZE_TITLE), ("Heading 4", SIZE_TITLE)):
        style = doc.styles[name]
        style.font.name = BODY_FONT
        style.font.size = Pt(size)
        style.font.bold = False
        style.font.italic = False
        style.font.color.rgb = BLACK

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = BODY_FONT
        style.font.size = Pt(SIZE_BODY)
        style.font.bold = False
        style.font.color.rgb = BLACK


def is_separator(row):
    """True for the | --- | --- | line that separates a Markdown table header."""
    return all(cell for cell in row) and all(set(cell) <= set("-: ") for cell in row)


def markdown_to_docx(lines, doc):
    in_code = False
    code_buffer = []
    table_buffer = []

    def flush_table():
        if table_buffer:
            add_table(doc, [row for row in table_buffer if not is_separator(row)])
            table_buffer.clear()

    for raw in lines + ["", ""]:
        line = raw.rstrip()

        if line.strip().startswith("```"):
            if in_code:
                add_code_block(doc, code_buffer)
                code_buffer = []
            else:
                flush_table()
            in_code = not in_code
            continue

        if in_code:
            code_buffer.append(line)
            continue

        if line.strip().startswith("|"):
            table_buffer.append([cell.strip() for cell in line.strip().strip("|").split("|")])
            continue
        flush_table()

        stripped = line.strip()
        if not stripped:
            continue
        if stripped.startswith("# "):
            add_heading(doc, stripped[2:], SIZE_HEADING)
        elif stripped.startswith("#### "):
            add_heading(doc, stripped[5:], SIZE_TITLE)
        elif stripped.startswith("### "):
            add_heading(doc, stripped[4:], SIZE_SUBHEADING)
        elif stripped.startswith("## "):
            add_heading(doc, stripped[3:], SIZE_HEADING)
        elif stripped.startswith("- "):
            add_point(doc, stripped[2:], "List Bullet")
        elif NUMBERED_RE.match(stripped):
            add_point(doc, NUMBERED_RE.match(stripped).group(1), "List Number")
        else:
            add_body(doc, stripped)


def main():
    if not SOURCE.exists():
        print(f"ERROR: {SOURCE} not found")
        return 2

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.9)
    section.bottom_margin = Inches(0.9)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    configure_styles(doc)

    markdown_to_docx(SOURCE.read_text(encoding="utf-8").splitlines(), doc)

    paragraphs = len(doc.paragraphs)
    tables = len(doc.tables)
    doc.save(OUTPUT)
    print(f"Wrote {OUTPUT.name}")
    print(f"  paragraphs: {paragraphs}   tables: {tables}   size: {OUTPUT.stat().st_size // 1024} KB")
    return 0


if __name__ == "__main__":
    sys.exit(main())
